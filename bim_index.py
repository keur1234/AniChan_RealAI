"""Index of every index code in the BIM exports, so "is this in BIM?" is answered without downloading files.

Each BIM export (one per component, e.g. CUP_25Digits_Combined_20261002.xlsx) has a sheet
Combined_25Digits with the index code in OBK.FM_25Digits_TAG. A background sync reads only the
files that changed since last time and stores the tags in a local SQLite database (bim_index.db).
Lookups query that database and take milliseconds.

Where the files come from (set one):
- BIM_LOCAL_FOLDER     a folder on this computer, e.g. the folder Google Drive for Desktop syncs
- BIM_DRIVE_FOLDER     the Google Drive folder link/ID, read with the Drive API using
                       GOOGLE_SERVICE_ACCOUNT_FILE (folder shared with the service account)
                       or BIM_DRIVE_API_KEY (folder shared as "Anyone with the link")
Only the newest file per component is used (by the date in the file name, then modified time).
"""
import fnmatch
import io
import logging
import os
import re
import sqlite3
import threading
import time
from datetime import datetime, timedelta, timezone

log = logging.getLogger("vicky")
APP_DIR = os.path.dirname(os.path.abspath(__file__))
BANGKOK = timezone(timedelta(hours=7))

DB_PATH = os.getenv("BIM_DB") or os.path.join(APP_DIR, "bim_index.db")
LOCAL_FOLDER = os.getenv("BIM_LOCAL_FOLDER", "")
DRIVE_FOLDER = os.getenv("BIM_DRIVE_FOLDER", "")
SERVICE_ACCOUNT_FILE = os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE", "")
DRIVE_API_KEY = os.getenv("BIM_DRIVE_API_KEY", "")
FILE_PATTERN = os.getenv("BIM_FILE_PATTERN", "*_25Digits_Combined*.xlsx")
SHEET = os.getenv("BIM_SHEET", "Combined_25Digits")
TAG_COLUMN = os.getenv("BIM_TAG_COLUMN", "OBK.FM_25Digits_TAG")
SYNC_SECONDS = float(os.getenv("BIM_SYNC_HOURS", "12")) * 3600

# Columns kept for each element (missing ones are stored empty)
DETAIL_COLUMNS = {
    "discipline": "Discipline", "category": "Category", "family": "Family", "type": "Type",
    "guid": "GUID", "element_id": "Element ID", "source_file": "Source File",
    "source_sheet": "Source Sheet", "status": "Status",
}

_lock = threading.Lock()
_state = {"syncing": False, "last_attempt": 0.0, "last_error": ""}


# ---------------------------------------------------------------- database
def _connect():
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.execute("""CREATE TABLE IF NOT EXISTS files (
        key TEXT PRIMARY KEY, component TEXT, name TEXT, version TEXT, source_id TEXT,
        modified TEXT, link TEXT, rows INTEGER, synced_at TEXT)""")
    con.execute(f"""CREATE TABLE IF NOT EXISTS elements (
        tag TEXT, file_key TEXT, {', '.join(f'{c} TEXT' for c in DETAIL_COLUMNS)})""")
    con.execute("CREATE INDEX IF NOT EXISTS idx_tag ON elements(tag)")
    con.execute("CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT)")
    return con


def _now():
    return datetime.now(BANGKOK).strftime("%Y-%m-%d %H:%M")


def normalize(code):
    return str(code).strip().upper()


# ---------------------------------------------------------------- choosing files
def _component_and_version(name):
    """CUP_25Digits_Combined_20261002.xlsx -> ('CUP', '20261002')."""
    stem = os.path.splitext(os.path.basename(name))[0]
    component = re.split(r"_25Digits", stem, flags=re.IGNORECASE)[0] or stem
    dates = re.findall(r"(\d{8})", stem)
    return component.upper(), (dates[-1] if dates else "")


def _latest_per_component(files):
    """Keep the newest file for each component."""
    best = {}
    for f in files:
        component, version = _component_and_version(f["name"])
        f["component"], f["version"] = component, version
        current = best.get(component)
        if current is None or (version, f["modified"]) > (current["version"], current["modified"]):
            best[component] = f
    return list(best.values())


# ---------------------------------------------------------------- sources
def _local_files():
    files = []
    for root, _, names in os.walk(LOCAL_FOLDER):
        for name in names:
            if fnmatch.fnmatch(name.lower(), FILE_PATTERN.lower()) and not name.startswith("~$"):
                path = os.path.join(root, name)
                files.append({"name": name, "source_id": path, "link": path,
                              "modified": datetime.fromtimestamp(os.path.getmtime(path), BANGKOK).isoformat()})
    return files


def _local_read(f):
    with open(f["source_id"], "rb") as fh:
        return fh.read()


def _folder_id(value):
    m = re.search(r"/folders/([A-Za-z0-9_-]{10,})", value) or re.search(r"[?&]id=([A-Za-z0-9_-]{10,})", value)
    return m.group(1) if m else value.strip()


def _drive_auth():
    """(headers, params) for Drive API calls."""
    if SERVICE_ACCOUNT_FILE:
        from google.auth.transport.requests import Request
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_file(
            SERVICE_ACCOUNT_FILE, scopes=["https://www.googleapis.com/auth/drive.readonly"])
        creds.refresh(Request())
        return {"Authorization": f"Bearer {creds.token}"}, {}
    if DRIVE_API_KEY:
        return {}, {"key": DRIVE_API_KEY}
    raise RuntimeError("set GOOGLE_SERVICE_ACCOUNT_FILE or BIM_DRIVE_API_KEY to read the Drive folder")


def _drive_files():
    import requests
    headers, params = _drive_auth()
    files, folders = [], [_folder_id(DRIVE_FOLDER)]
    while folders:                                         # walk sub-folders too
        folder = folders.pop()
        page = None
        while True:
            q = {"q": f"'{folder}' in parents and trashed=false", "pageSize": 1000,
                 "fields": "nextPageToken,files(id,name,mimeType,modifiedTime,webViewLink)",
                 "supportsAllDrives": "true", "includeItemsFromAllDrives": "true", **params}
            if page:
                q["pageToken"] = page
            r = requests.get("https://www.googleapis.com/drive/v3/files", headers=headers, params=q, timeout=60)
            r.raise_for_status()
            data = r.json()
            for item in data.get("files", []):
                if item["mimeType"] == "application/vnd.google-apps.folder":
                    folders.append(item["id"])
                    continue
                name = item["name"]
                is_sheet = item["mimeType"] == "application/vnd.google-apps.spreadsheet"
                if is_sheet:
                    name += ".xlsx"
                if fnmatch.fnmatch(name.lower(), FILE_PATTERN.lower()):
                    files.append({"name": name, "source_id": item["id"], "modified": item["modifiedTime"],
                                  "link": item.get("webViewLink", ""), "google_sheet": is_sheet})
            page = data.get("nextPageToken")
            if not page:
                break
    return files


def _drive_read(f):
    import requests
    headers, params = _drive_auth()
    base = f"https://www.googleapis.com/drive/v3/files/{f['source_id']}"
    if f.get("google_sheet"):
        url, extra = base + "/export", {"mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    else:
        url, extra = base, {"alt": "media", "supportsAllDrives": "true"}
    r = requests.get(url, headers=headers, params={**extra, **params}, timeout=300)
    r.raise_for_status()
    return r.content


def configured():
    return bool(LOCAL_FOLDER or DRIVE_FOLDER)


# ---------------------------------------------------------------- sync
def _read_tags(data):
    import pandas as pd
    df = pd.read_excel(io.BytesIO(data), sheet_name=SHEET, dtype=str)
    if TAG_COLUMN not in df.columns:
        raise RuntimeError(f"column {TAG_COLUMN} not found in sheet {SHEET}")
    df = df[df[TAG_COLUMN].notna()]
    rows = []
    for rec in df.to_dict("records"):
        tag = normalize(rec[TAG_COLUMN])
        if tag and tag != "NAN":
            rows.append((tag, *[("" if pd.isna(rec.get(col)) else str(rec.get(col))) for col in DETAIL_COLUMNS.values()]))
    return rows


def sync(force=False):
    """Index new or changed BIM files. Returns a short report."""
    if not configured():
        return "BIM source not configured (set BIM_LOCAL_FOLDER or BIM_DRIVE_FOLDER)"
    with _lock:
        _state["last_attempt"] = time.time()
        use_drive = bool(DRIVE_FOLDER) and not LOCAL_FOLDER
        try:
            listed = _drive_files() if use_drive else _local_files()
        except Exception as e:
            _state["last_error"] = f"listing failed: {e}"
            log.warning("BIM sync: %s", _state["last_error"])
            return _state["last_error"]
        wanted = _latest_per_component(listed)
        con = _connect()
        try:
            known = {r[0]: r for r in con.execute("SELECT key, source_id, modified FROM files")}
            updated, errors = [], []
            for f in wanted:
                key = f["component"]
                if not force and key in known and known[key][1] == f["source_id"] and known[key][2] == f["modified"]:
                    continue
                try:
                    rows = _read_tags(_drive_read(f) if use_drive else _local_read(f))
                except Exception as e:
                    errors.append(f"{f['name']}: {e}")
                    log.warning("BIM sync: %s failed: %s", f["name"], e)
                    continue
                with con:                                  # replace this component atomically
                    con.execute("DELETE FROM elements WHERE file_key=?", (key,))
                    con.executemany(f"INSERT INTO elements VALUES ({', '.join('?' * (len(DETAIL_COLUMNS) + 2))})",
                                    [(r[0], key, *r[1:]) for r in rows])
                    con.execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?,?,?,?,?,?)",
                                (key, f["component"], f["name"], f["version"], f["source_id"], f["modified"],
                                 f["link"], len(rows), _now()))
                updated.append(f"{f['name']} {len(rows):,}")
                log.info("BIM sync: indexed %s (%d tags)", f["name"], len(rows))
            # components that disappeared from the source
            gone = set(known) - {f["component"] for f in wanted}
            with con:
                for key in gone:
                    con.execute("DELETE FROM elements WHERE file_key=?", (key,))
                    con.execute("DELETE FROM files WHERE key=?", (key,))
                con.execute("INSERT OR REPLACE INTO meta VALUES ('last_sync', ?)", (_now(),))
        finally:
            con.close()
        _state["last_error"] = "; ".join(errors)
        parts = [f"updated {len(updated)} file(s)" + (": " + ", ".join(updated) if updated else ""),
                 f"unchanged {len(wanted) - len(updated) - len(errors)}"]
        if gone:
            parts.append(f"removed {', '.join(sorted(gone))}")
        if errors:
            parts.append("errors: " + "; ".join(errors))
        return "; ".join(parts)


def sync_in_background(force=False):
    if _state["syncing"]:
        return False

    def run():
        try:
            sync(force)
        finally:
            _state["syncing"] = False
    _state["syncing"] = True
    threading.Thread(target=run, daemon=True).start()
    return True


def maybe_sync():
    """Start a background sync when the index is older than BIM_SYNC_HOURS."""
    if configured() and time.time() - _state["last_attempt"] >= SYNC_SECONDS:
        sync_in_background()


# ---------------------------------------------------------------- queries
def last_sync():
    if not os.path.exists(DB_PATH):
        return None
    con = _connect()
    try:
        row = con.execute("SELECT v FROM meta WHERE k='last_sync'").fetchone()
        return row[0] if row else None
    finally:
        con.close()


def lookup(codes, per_code=20):
    """{code: [element dicts]} for each code (empty list = not in BIM)."""
    codes = [normalize(c) for c in codes if str(c).strip()]
    out = {c: [] for c in codes}
    if not codes or not os.path.exists(DB_PATH):
        return out
    con = _connect()
    con.row_factory = sqlite3.Row
    try:
        cols = ", ".join(f"e.{c}" for c in DETAIL_COLUMNS)
        for i in range(0, len(codes), 500):
            chunk = list(dict.fromkeys(codes[i:i + 500]))
            q = (f"SELECT e.tag, e.file_key, f.name AS bim_file, f.link, {cols} FROM elements e "
                 f"JOIN files f ON f.key = e.file_key WHERE e.tag IN ({','.join('?' * len(chunk))})")
            for row in con.execute(q, chunk):
                if len(out[row["tag"]]) < per_code:
                    out[row["tag"]].append(dict(row))
        return out
    finally:
        con.close()


def count_by_tag(codes):
    """{code: number of BIM elements} for many codes at once."""
    codes = list(dict.fromkeys(normalize(c) for c in codes if str(c).strip()))
    out = dict.fromkeys(codes, 0)
    if not codes or not os.path.exists(DB_PATH):
        return out
    con = _connect()
    try:
        for i in range(0, len(codes), 500):
            chunk = codes[i:i + 500]
            for tag, n in con.execute(f"SELECT tag, COUNT(*) FROM elements WHERE tag IN ({','.join('?' * len(chunk))}) "
                                      "GROUP BY tag", chunk):
                out[tag] = n
        return out
    finally:
        con.close()


def status():
    """Files in the index, for /bimstatus."""
    info = {"configured": configured(), "source": "local folder" if LOCAL_FOLDER else ("Google Drive" if DRIVE_FOLDER else "-"),
            "last_sync": last_sync(), "syncing": _state["syncing"], "error": _state["last_error"], "files": []}
    if os.path.exists(DB_PATH):
        con = _connect()
        try:
            info["files"] = [{"component": c, "name": n, "rows": r} for c, n, r in
                             con.execute("SELECT component, name, rows FROM files ORDER BY component")]
        finally:
            con.close()
    return info


# ---------------------------------------------------------------- reports
def _t(lang, th, en):
    return th if lang == "th" else en


def _pct(n, total):
    return f"{(n / total * 100 if total else 0):.2f}%"


def _as_of(lang):
    when = last_sync()
    return _t(lang, f"ข้อมูล BIM ณ {when}" if when else "ยังไม่มีข้อมูล BIM",
              f"BIM data as of {when}" if when else "No BIM data yet")


def _element_line(e):
    parts = [e["bim_file"], e["discipline"], e["category"], e["family"]]
    return " | ".join(p for p in parts if p) + (f" | Status {e['status']}" if e.get("status") else "")


def format_lookup(results, lang="th", details_per_code=3):
    total = len(results)
    found = [c for c, els in results.items() if els]
    lines = [_t(lang, "ผลค้น index codes ใน BIM", "Index codes BIM lookup"), _as_of(lang),
             _t(lang, f"index codes ทั้งหมด: {total:,}", f"Total index codes: {total:,}"),
             _t(lang, f"พบใน BIM: {len(found):,} index codes {_pct(len(found), total)}",
                f"Found in BIM: {len(found):,} index codes {_pct(len(found), total)}"),
             _t(lang, f"ไม่พบใน BIM: {total - len(found):,} index codes {_pct(total - len(found), total)}",
                f"Not in BIM: {total - len(found):,} index codes {_pct(total - len(found), total)}"), ""]
    for i, (code, els) in enumerate(results.items(), 1):
        if els:
            lines.append(f"{i}. {code}: " + _t(lang, f"พบ {len(els)} elements", f"{len(els)} elements"))
            lines += [f"   - {_element_line(e)}" for e in els[:details_per_code]]
            if len(els) > details_per_code:
                lines.append(_t(lang, f"   และอีก {len(els) - details_per_code} elements", f"   and {len(els) - details_per_code} more"))
        else:
            lines.append(f"{i}. {code}: " + _t(lang, "ไม่พบใน BIM", "not in BIM"))
    return "\n".join(lines).strip()


def _codes_from_file(path):
    """(DataFrame of 'index codes' + 'Source Sheet', column used) using the validator's own file reader."""
    import validate_index_code as vic
    log_lines = []
    df = vic.read_input(path, log_lines)
    if df is None:
        return None, None
    df = df.loc[:, ~df.columns.duplicated(keep="first")]
    col = vic.find_check_column(df)
    if col is None:
        return None, None
    df = df.dropna(subset=[col])
    codes = df[col].astype(str).map(lambda x: x if "_" not in x else "_".join(x.split("_")[:-1])).map(normalize)
    out = df.assign(**{"index codes": codes})[["index codes", "Source Sheet"]]
    return out[(out["index codes"] != "") & (out["index codes"] != "NAN")], str(col)


def _tags_with_components(components):
    if not os.path.exists(DB_PATH) or not components:
        return set()
    con = _connect()
    try:
        tags = set()
        for comp in components:
            tags |= {r[0] for r in con.execute("SELECT DISTINCT tag FROM elements WHERE tag LIKE ?", (f"{comp}-%",))}
        return tags
    finally:
        con.close()


def compare_file(path, file_name, out_dir, lang="th"):
    """Check every index code of a file against BIM. Returns (reply text, xlsx path or None)."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    codes_df, col = _codes_from_file(path)
    if codes_df is None or codes_df.empty:
        return _t(lang, f"หนูไม่พบคอลัมน์ index codes ในไฟล์ {file_name} ค่ะ",
                  f"I couldn't find an index codes column in {file_name}."), None

    unique = list(dict.fromkeys(codes_df["index codes"]))
    details = lookup(unique, per_code=1)
    counts = count_by_tag(unique)
    codes_df = codes_df.assign(**{
        "Found in BIM": codes_df["index codes"].map(lambda c: "Yes" if counts.get(c) else "No"),
        "BIM elements": codes_df["index codes"].map(lambda c: counts.get(c, 0)),
        "BIM file": codes_df["index codes"].map(lambda c: details[c][0]["bim_file"] if details.get(c) else ""),
        "Discipline": codes_df["index codes"].map(lambda c: details[c][0]["discipline"] if details.get(c) else ""),
        "Category": codes_df["index codes"].map(lambda c: details[c][0]["category"] if details.get(c) else ""),
        "Family": codes_df["index codes"].map(lambda c: details[c][0]["family"] if details.get(c) else ""),
        "BIM Status": codes_df["index codes"].map(lambda c: details[c][0]["status"] if details.get(c) else ""),
    })
    total = len(unique)
    found = sum(1 for c in unique if counts.get(c))
    missing = [c for c in unique if not counts.get(c)]
    components = {c.split("-")[0] for c in unique if "-" in c}
    only_in_bim = sorted(_tags_with_components(components) - set(unique))

    stem = os.path.splitext(os.path.basename(file_name))[0]
    path_out = os.path.join(out_dir, f"{stem}_BIM_check.xlsx")
    summary = pd.DataFrame([
        ["File", file_name], ["Column checked", col], ["BIM data as of", last_sync() or "-"],
        ["Total index codes", total], ["Found in BIM", found], ["Found in BIM %", found / total if total else 0],
        ["Not in BIM", len(missing)], ["In BIM but not in file", len(only_in_bim)],
        ["Components compared", ", ".join(sorted(components))],
    ])
    sheets = [("BIM Check", codes_df), ("Not in BIM", pd.DataFrame({"index codes": missing})),
              ("In BIM not in file", pd.DataFrame({"index codes": only_in_bim}))]
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    with pd.ExcelWriter(path_out, engine="openpyxl") as w:
        summary.to_excel(w, sheet_name="Summary", index=False, header=False)
        w.sheets["Summary"]["B6"].number_format = "0.00%"
        for name, data in sheets:
            data.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "A2"
            if data.shape[1] and len(data):
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
        for ws in w.book.worksheets:
            for column in ws.columns:
                width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 60)

    lines = [_t(lang, "ผลเทียบ index codes กับ BIM", "Index codes vs BIM"), _t(lang, "ไฟล์: ", "File: ") + file_name,
             _as_of(lang),
             _t(lang, f"index codes ทั้งหมด: {total:,}", f"Total index codes: {total:,}"),
             _t(lang, f"พบใน BIM: {found:,} index codes {_pct(found, total)}",
                f"Found in BIM: {found:,} index codes {_pct(found, total)}"),
             _t(lang, f"ไม่พบใน BIM: {len(missing):,} index codes {_pct(len(missing), total)}",
                f"Not in BIM: {len(missing):,} index codes {_pct(len(missing), total)}")]
    lines += [f"  {c}" for c in missing[:10]]
    if len(missing) > 10:
        lines.append(_t(lang, f"  และอีก {len(missing) - 10:,} index codes ดูในไฟล์", f"  and {len(missing) - 10:,} more in the file"))
    lines.append(_t(lang, f"มีใน BIM แต่ไม่มีในไฟล์: {len(only_in_bim):,} index codes Component {', '.join(sorted(components))}",
                    f"In BIM but not in the file: {len(only_in_bim):,} index codes, components {', '.join(sorted(components))}"))
    return "\n".join(lines), path_out


def format_status(lang="th"):
    s = status()
    if not s["configured"]:
        return _t(lang, "ยังไม่ได้ตั้งค่าแหล่งไฟล์ BIM (BIM_LOCAL_FOLDER หรือ BIM_DRIVE_FOLDER)",
                  "BIM source not configured (BIM_LOCAL_FOLDER or BIM_DRIVE_FOLDER)")
    lines = [_t(lang, "สถานะข้อมูล BIM", "BIM index status"), f"Source: {s['source']}", _as_of(lang)]
    if s["syncing"]:
        lines.append(_t(lang, "กำลังซิงก์อยู่", "Sync in progress"))
    total = sum(f["rows"] for f in s["files"])
    lines.append(_t(lang, f"ไฟล์: {len(s['files'])} Component, {total:,} elements", f"Files: {len(s['files'])} components, {total:,} elements"))
    lines += [f"- {f['component']}: {f['name']} {f['rows']:,}" for f in s["files"]]
    if s["error"]:
        lines.append("Error: " + s["error"])
    return "\n".join(lines)

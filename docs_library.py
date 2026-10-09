"""Team documents Vicky can hand out by keyword (e.g. "Raw file summary" -> Google Drive link).

The list lives in documents.xlsx next to the code (not committed: the links are internal).
Columns: Keyword | Link | Other keywords (comma-separated) | Note.
Edit it in Excel and save; Vicky picks up the changes on the next message.
"""
import difflib
import logging
import os
import re
import threading

DOCUMENTS_FILE = os.getenv("DOCUMENTS_FILE") or os.path.join(os.path.dirname(os.path.abspath(__file__)), "documents.xlsx")
HEADERS = ["Keyword", "Link", "Other keywords", "Note"]

log = logging.getLogger("vicky")
_lock = threading.Lock()
_cache = {"mtime": None, "docs": []}


def _norm(text):
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def load():
    """Documents from the Excel file, re-read only when the file changes."""
    try:
        mtime = os.path.getmtime(DOCUMENTS_FILE)
    except OSError:
        return []
    with _lock:
        if _cache["mtime"] == mtime:
            return _cache["docs"]
        try:
            from openpyxl import load_workbook
            wb = load_workbook(DOCUMENTS_FILE, read_only=True, data_only=True)
            rows = list(wb.worksheets[0].iter_rows(values_only=True))
            wb.close()
        except Exception as e:                     # e.g. half-saved file; keep the previous list
            log.warning("Could not read %s: %s", DOCUMENTS_FILE, e)
            return _cache["docs"]
        docs = []
        for row in rows[1:]:
            cells = [("" if c is None else str(c).strip()) for c in (list(row) + [None] * 4)[:4]]
            keyword, link, others, note = cells
            if keyword and link.startswith("http"):
                docs.append({"name": keyword, "url": link, "note": note,
                             "aliases": [a.strip() for a in re.split(r"[,\n]", others) if a.strip()]})
        _cache.update(mtime=mtime, docs=docs)
        log.info("Loaded %d documents from %s", len(docs), DOCUMENTS_FILE)
        return docs


def _keys(doc):
    return [_norm(k) for k in [doc["name"], *doc.get("aliases", [])] if k]


def mentioned_in(text):
    """Documents whose keyword appears in the text, longest match first."""
    text = _norm(text)
    hits = []
    for doc in load():
        best = max((len(k) for k in _keys(doc) if k in text), default=0)
        if best:
            hits.append((best, doc))
    return [d for _, d in sorted(hits, key=lambda h: -h[0])]


# Words that say nothing about which document is meant
GENERIC_WORDS = ("ตาราง", "ไฟล์", "เอกสาร", "ขอ", "หน่อย", "ด้วย", "ที", "ครับ", "ค่ะ", "คะ", "ลิงก์", "ลิ้งค์", "ลิ้ง",
                 "table", "file", "document", "doc", "sheet", "link", "please", "the")


def _strip_generic(text):
    text = _norm(text)
    for w in GENERIC_WORDS:
        text = text.replace(w, " ")
    return re.sub(r"[\s/|,.]+", " ", text).strip()


def search(query):
    """(matches, similar): matches are documents the query clearly names; similar are only close guesses."""
    found = mentioned_in(query)
    if found:
        return found, []
    q = _strip_generic(query)
    if not q:
        return [], []
    strong, similar = [], []
    for doc in load():
        keys = [_strip_generic(k) for k in [doc["name"], *doc.get("aliases", [])]]
        keys = [k for k in keys if k]
        score = max((difflib.SequenceMatcher(None, q, k).ratio() for k in keys), default=0)
        if len(q) >= 6 and any(q in k or k in q for k in keys if len(k) >= 6):
            score = max(score, 0.9)
        if score >= 0.85:
            strong.append((score, doc))
        elif score >= 0.6:
            similar.append((score, doc))
    by_score = lambda items: [d for _, d in sorted(items, key=lambda s: -s[0])]
    return by_score(strong), by_score(similar)


class DocumentAccessError(Exception):
    pass


def download(doc, timeout=60):
    """Download a team document (Google Drive / Sheets link) as an Excel file. Returns (file_name, bytes).

    Works when the file is shared as "Anyone with the link can view".
    """
    import requests
    url = doc["url"]
    m = re.search(r"/d/([A-Za-z0-9_-]{20,})", url) or re.search(r"[?&]id=([A-Za-z0-9_-]{20,})", url)
    if not m:
        raise DocumentAccessError("unsupported link; only Google Drive / Google Sheets links can be opened")
    file_id = m.group(1)
    candidates = [f"https://drive.google.com/uc?export=download&id={file_id}"]
    if "/spreadsheets/" in url:
        candidates.insert(0 if "rtpof=true" not in url else 1,
                          f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx")
    for candidate in candidates:
        try:
            r = requests.get(candidate, timeout=timeout, allow_redirects=True)
        except requests.exceptions.RequestException as e:
            log.warning("Download of %s failed: %s", doc["name"], e)
            continue
        if r.status_code == 200 and r.content[:2] == b"PK":            # .xlsx files are zip archives
            name = re.sub(r'[\\/:*?"<>|]', "_", doc["name"]) + ".xlsx"
            return name, r.content
    raise DocumentAccessError("the file could not be downloaded; it must be shared as "
                              "'Anyone with the link can view' on Google Drive")


def link_lines(docs):
    return [f"- {d['name']}\n{d['url']}" for d in docs]


def write_template(path, rows=()):
    """Create a documents.xlsx with the expected columns (and optional rows)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append(list(row))
    for col, width in zip("ABCD", (28, 70, 45, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)

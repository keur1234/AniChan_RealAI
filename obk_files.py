"""Validate Excel/CSV files sent in LINE and keep the outputs downloadable for a while."""
import os
import re
import secrets
import shutil
import tempfile
import time

import obk_validator

ALLOWED_EXTENSIONS = ('.xlsx', '.xlsm', '.xls', '.csv')
MAX_FILE_BYTES = int(os.getenv("OBK_MAX_FILE_MB", "20")) * 1024 * 1024
DOWNLOAD_TTL_SECONDS = int(os.getenv("OBK_DOWNLOAD_TTL_HOURS", "24")) * 3600
WORK_ROOT = os.path.join(tempfile.gettempdir(), "anichan_obk")

# Thai fonts to fall back to for Thai file/sheet names in the pie chart (Windows, Linux, macOS)
THAI_FONTS = ['Leelawadee UI', 'Leelawadee', 'Tahoma', 'Loma', 'Garuda', 'Noto Sans Thai', 'Thonburi']

try:
    import matplotlib
    from matplotlib import font_manager
    installed = {f.name for f in font_manager.fontManager.ttflist}
    matplotlib.rcParams['font.family'] = ['DejaVu Sans'] + [f for f in THAI_FONTS if f in installed]
except ImportError:
    pass

# Each job lives in WORK_ROOT/<random token>/ on disk, so download links keep
# working after the bot restarts. Jobs older than the TTL are deleted.
TOKEN_RE = re.compile(r'^[A-Za-z0-9_-]{16,64}$')


def is_supported(file_name):
    return file_name.lower().endswith(ALLOWED_EXTENSIONS)


def safe_name(file_name):
    name = os.path.basename(file_name or "input.xlsx")
    return re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', name) or "input.xlsx"


def _cleanup_expired():
    if not os.path.isdir(WORK_ROOT):
        return
    cutoff = time.time() - DOWNLOAD_TTL_SECONDS
    for token in os.listdir(WORK_ROOT):
        job_dir = os.path.join(WORK_ROOT, token)
        try:
            if os.path.getmtime(job_dir) < cutoff:
                shutil.rmtree(job_dir, ignore_errors=True)
        except OSError:
            pass


def get_download(token, file_name):
    """Path of an output file of a job, or None if it doesn't exist or has expired."""
    if not TOKEN_RE.match(token) or os.path.basename(file_name) != file_name:
        return None
    out_dir = os.path.join(WORK_ROOT, token, "out")
    if not os.path.isdir(out_dir) or os.path.getmtime(os.path.join(WORK_ROOT, token)) < time.time() - DOWNLOAD_TTL_SECONDS:
        return None
    for root, _, files in os.walk(out_dir):
        if file_name in files:
            return os.path.join(root, file_name)
    return None


def validate_file(file_name, data):
    """Run the original file pipeline. Returns (result dict, token for downloads)."""
    # Imported lazily so the chat bot still starts if pandas isn't installed
    import validate_index_code as vic

    _cleanup_expired()
    token = secrets.token_urlsafe(16)
    job_dir = os.path.join(WORK_ROOT, token)
    os.makedirs(job_dir, exist_ok=True)
    path = os.path.join(job_dir, safe_name(file_name))
    with open(path, 'wb') as f:
        f.write(data)

    result = vic.process_file(path, obk_validator.get_master(), os.path.join(job_dir, "out"), room=True, dup_mode='sheet')
    return result, token


SKIP_REASONS_EN = {
    'ไฟล์ว่าง': 'the file is empty',
    'ไม่พบคอลัมน์ Asset ID': 'no index code column was found',
    'ไม่มี Asset ID ที่ใช้ได้': 'no usable index codes were found',
}


def format_summary(result, file_name, lang='th'):
    """File summary in the agreed OBK format: total, Incorrect, Validation result."""
    t = lambda th, en: obk_validator._t(lang, th, en)
    if result.get('skipped'):
        reason = (result['skipped'].replace('ไม่พบคอลัมน์ Asset ID', 'ไม่พบคอลัมน์ index codes')
                  .replace('ไม่มี Asset ID ที่ใช้ได้', 'ไม่มี index codes ที่ใช้ได้')
                  if lang == 'th' else SKIP_REASONS_EN.get(result['skipped'], result['skipped']))
        msg = [t(f"หนูไม่สามารถ validate ไฟล์ {file_name} ได้ค่ะ: {reason}", f"I couldn't validate {file_name}: {reason}")]
        if result.get('columns'):
            msg.append(t("คอลัมน์ที่พบในไฟล์: ", "Columns found: ") + ", ".join(result['columns'][:15]))
            msg.append(t("หนูอ่าน index codes จากคอลัมน์ชื่อ Asset ID, IndexCode, Index, Equipment ID หรือ Name "
                         "กรุณาแก้ชื่อหัวคอลัมน์แล้วส่งไฟล์อีกครั้งค่ะ",
                         "I read index codes from columns named Asset ID, IndexCode, Index, Equipment ID or Name. "
                         "Please fix the header row and send the file again."))
        return "\n".join(msg)

    type_counts = {typ: v['count'] for typ, v in result['types'].items()}
    lines = [t("ผล validation index codes", "Index codes validation result"), t("ไฟล์: ", "File: ") + file_name]
    dup_lines = [t(f"  {code} ซ้ำ {n} ครั้ง", f"  {code} appears {n} times") for code, n in result.get('duplicate_codes', [])]
    lines += obk_validator.summary_lines(result['records'], result.get('duplicate_codes_total', 0), dup_lines,
                                         type_counts.get('TYPE B', 0), type_counts, result.get('type_c_codes', []), lang)
    return "\n".join(lines)


TYPE_COLUMNS = ['TYPE A', 'TYPE B', 'TYPE B OR C', 'TYPE C', 'N/A']


def _reason_display(text, lang):
    return "; ".join(obk_validator.reason_text(r) for r in str(text).split('; ') if r and r != 'nan')


def _pivot(df, by):
    """index codes per TYPE for each value of `by`, plus totals and TYPE A share."""
    import pandas as pd
    table = pd.crosstab(df[by].fillna('-').astype(str), df['TYPE'])
    types = [t for t in TYPE_COLUMNS if t in table.columns] + [t for t in table.columns if t not in TYPE_COLUMNS]
    table = table.reindex(columns=types, fill_value=0)
    table['Total index codes'] = table.sum(axis=1)
    table['TYPE A %'] = (table.get('TYPE A', 0) / table['Total index codes']).fillna(0)
    return table.reset_index().rename(columns={by: by if by != 'Asset Category Full name' else 'Asset Category'})


def build_summary_workbook(result, file_name, lang='th'):
    """Excel summary table of a validated raw data file. Returns its path (next to the other outputs)."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    main_path = next(p for p in result['outputs'] if p.endswith('_output.xlsx'))
    df = pd.read_excel(main_path, sheet_name='Sheet1')
    total = len(df)
    type_counts = df['TYPE'].value_counts()
    dup = df[df['Dup Count'].fillna(0) > 1].drop_duplicates('Asset ID')
    reason_col = 'validation_result_thai' if lang == 'th' else 'validation_result'

    summary_rows = [
        ['File', file_name],
        ['Total index codes', total],
        ['Incorrect - TYPE C', int(type_counts.get('TYPE C', 0))],
        ['Incorrect - duplicate index codes', len(dup)],
        ['Incorrect - duplicate running numbers TYPE B', int(type_counts.get('TYPE B', 0))],
    ]
    validation_rows = [[t, int(type_counts.get(t, 0)), type_counts.get(t, 0) / total if total else 0]
                       for t in TYPE_COLUMNS + [t for t in type_counts.index if t not in TYPE_COLUMNS]
                       if type_counts.get(t, 0)]

    incorrect = df[df['TYPE'] == 'TYPE C'].drop_duplicates('Asset ID')
    incorrect = pd.DataFrame({
        'index codes': incorrect['Asset ID'],
        'Source Sheet': incorrect['Source Sheet'],
        'Rules': incorrect['Rules'],
        'Reason': incorrect[reason_col].map(lambda r: _reason_display(r, lang)),
    })
    duplicates = pd.DataFrame({
        'index codes': dup['Asset ID'],
        'Times found': dup['Dup Count'].astype(int),
        'Source Sheets': dup['Source Sheets'],
    }).sort_values('Times found', ascending=False)

    sheets = [
        ('By Sheet', _pivot(df, 'Source Sheet')),
        ('By Category', _pivot(df, 'Asset Category Full name')),
        ('By Component', _pivot(df, 'Component')),
        ('Incorrect TYPE C', incorrect),
        ('Duplicates', duplicates),
    ]

    stem = os.path.splitext(os.path.basename(file_name))[0]
    path = os.path.join(os.path.dirname(main_path), f"{stem}_summary.xlsx")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")

    with pd.ExcelWriter(path, engine='openpyxl') as w:
        # Summary sheet: key figures, then the Validation result table
        pd.DataFrame(summary_rows).to_excel(w, sheet_name='Summary', index=False, header=False)
        start = len(summary_rows) + 1
        pd.DataFrame(validation_rows, columns=['Validation result', 'index codes', 'Share']).to_excel(
            w, sheet_name='Summary', index=False, startrow=start)
        ws = w.sheets['Summary']
        for row in ws.iter_rows(min_row=1, max_row=len(summary_rows)):
            row[0].font = Font(bold=True)
        for cell in ws[start + 1]:
            cell.font, cell.fill = header_font, header_fill
        for row in ws.iter_rows(min_row=start + 2, max_row=start + 1 + len(validation_rows), min_col=3, max_col=3):
            row[0].number_format = '0.00%'

        for name, data in sheets:
            data.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = 'A2'
            if data.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            if 'TYPE A %' in data.columns:
                col = get_column_letter(list(data.columns).index('TYPE A %') + 1)
                for cell in ws[col][1:]:
                    cell.number_format = '0.00%'

        for ws in w.book.worksheets:
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
    return path


def new_job():
    """A fresh downloadable job folder. Returns (token, out_dir); files in out_dir are served at /files/<token>/<name>."""
    _cleanup_expired()
    token = secrets.token_urlsafe(16)
    out_dir = os.path.join(WORK_ROOT, token, "out")
    os.makedirs(out_dir, exist_ok=True)
    return token, out_dir

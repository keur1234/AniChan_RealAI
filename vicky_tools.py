"""Files Vicky can create on request, and the tool declarations Gemini sees."""
import os
import re

from google.genai import types

import obk_files

TEXT_EXTENSIONS = ('.csv', '.txt', '.md', '.json')
MAX_PREVIEW_ROWS = 50
MAX_CELL_CHARS = 120
NUMBER_RE = re.compile(r'^-?(0|[1-9]\d*)(\.\d+)?$')     # keeps codes like "001" as text


def _cell(value):
    """Turn numeric strings into numbers so Excel can sum them; leave everything else as text."""
    if isinstance(value, str) and NUMBER_RE.match(value.strip()):
        v = value.strip()
        return float(v) if '.' in v else int(v)
    return value


def _file_name(name, default_ext, allowed):
    name = obk_files.safe_name(name or f"file{default_ext}")
    stem, ext = os.path.splitext(name)
    if ext.lower() not in allowed:
        name = f"{stem or 'file'}{default_ext}"
    return name


def create_excel(out_dir, file_name, sheets):
    """Write an .xlsx with one sheet per entry in `sheets` ({name, columns, rows}). Returns the path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for i, sheet in enumerate(sheets or [{"name": "Sheet1", "columns": [], "rows": []}], 1):
        title = re.sub(r'[\[\]:*?/\\]', '_', str(sheet.get("name") or f"Sheet{i}"))[:31] or f"Sheet{i}"
        while title in used:
            title = f"{title[:28]}_{i}"
        used.add(title)
        ws = wb.create_sheet(title)
        columns = sheet.get("columns") or []
        if columns:
            ws.append([str(c) for c in columns])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            ws.freeze_panes = 'A2'
        for row in sheet.get("rows") or []:
            ws.append([_cell(v) for v in (row if isinstance(row, list) else [row])])
        if ws.max_row > 1 and columns:
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)

    path = os.path.join(out_dir, _file_name(file_name, '.xlsx', ('.xlsx',)))
    wb.save(path)
    return path


def create_text(out_dir, file_name, content):
    path = os.path.join(out_dir, _file_name(file_name, '.txt', TEXT_EXTENSIONS))
    # utf-8-sig so Excel opens Thai CSV correctly
    encoding = 'utf-8-sig' if path.lower().endswith('.csv') else 'utf-8'
    with open(path, 'w', encoding=encoding, newline='') as f:
        f.write(content or "")
    return path


def preview_file(path, file_name):
    """Sheets, columns, row counts and the first rows of a spreadsheet, for Gemini to read."""
    import pandas as pd
    if file_name.lower().endswith('.csv'):
        frames = {os.path.splitext(file_name)[0]: pd.read_csv(path, dtype=str)}
    else:
        frames = pd.read_excel(path, sheet_name=None, dtype=str)
    sheets = []
    for name, df in frames.items():
        df = df.dropna(how='all')
        sheets.append({
            "sheet": str(name),
            "rows": int(len(df)),
            "columns": [str(c) for c in df.columns],
            "first_rows": df.head(MAX_PREVIEW_ROWS).fillna("").astype(str)
                            .map(lambda v: v[:MAX_CELL_CHARS]).values.tolist(),
        })
    return {"file": file_name, "sheets": sheets,
            "note": f"Only the first {MAX_PREVIEW_ROWS} rows of each sheet are shown; 'rows' is the full count."}


def _decl(name, description, properties=None, required=()):
    return types.FunctionDeclaration(
        name=name, description=description,
        parameters_json_schema={"type": "object", "properties": properties or {}, "required": list(required)})


SHEETS_SCHEMA = {
    "type": "array",
    "description": "Sheets of the workbook",
    "items": {
        "type": "object",
        "properties": {
            "name": {"type": "string", "description": "Sheet name"},
            "columns": {"type": "array", "items": {"type": "string"}, "description": "Header row"},
            "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}},
                     "description": "Data rows; every cell as a string (numbers like \"12.5\" become numbers)"},
        },
        "required": ["name", "columns", "rows"],
    },
}

FUNCTION_DECLARATIONS = [
    _decl("find_document_links",
          "Find the team's existing documents (e.g. Raw file summary, Final Validation summary) by name and send their links. "
          "Call this first whenever the user asks for a document, table or file.",
          {"query": {"type": "string", "description": "What the user asked for, e.g. 'Final Validation summary'"}},
          ["query"]),
    _decl("create_excel_file",
          "Create an Excel .xlsx file the user asked for (a table, list, plan, comparison...) and send them a download link. "
          "Use only data the user gave or data from other tools; never invent figures.",
          {"file_name": {"type": "string", "description": "File name, e.g. budget_plan.xlsx"}, "sheets": SHEETS_SCHEMA},
          ["file_name", "sheets"]),
    _decl("create_text_file",
          "Create a .csv, .txt, .md or .json file the user asked for and send them a download link.",
          {"file_name": {"type": "string", "description": "File name with extension, e.g. notes.txt"},
           "content": {"type": "string", "description": "Full file content"}},
          ["file_name", "content"]),
    _decl("read_document",
          "Open a team document (from find_document_links) and read its sheets, columns, row counts and first rows, "
          "to summarise it or answer questions about it.",
          {"name": {"type": "string", "description": "Document name or keyword"}}, ["name"]),
    _decl("read_raw_data_file",
          "Read the latest Excel/CSV file sent in this chat (or the one the user replied to): sheets, columns, row counts "
          "and the first rows. Use it to answer questions about the file or to build a custom table from it."),
    _decl("validate_index_codes",
          "Run the OBK validation program on index codes typed by the user. Only when the user asks to validate/check them.",
          {"codes": {"type": "array", "items": {"type": "string"}, "description": "Index codes, max 10"}},
          ["codes"]),
    _decl("validate_raw_data_file",
          "Run the OBK validation program on the latest file sent in this chat. Only when the user asks to validate/check it. "
          "Returns the validation summary and download links to the result files. "
          "Set document_name to validate a team document instead.",
          {"document_name": {"type": "string", "description": "Team document to use instead of the file sent in the chat (optional)"}}),
    _decl("summarize_raw_data_file",
          "Validate the latest raw data file in this chat and create the standard OBK Excel summary table "
          "(by sheet, category, component, Incorrect TYPE C, duplicates). Use when the user asks for a summary table of the raw data file. "
          "Set document_name to use a team document instead.",
          {"document_name": {"type": "string", "description": "Team document to use instead of the file sent in the chat (optional)"}}),
]

FUNCTION_TOOL = types.Tool(function_declarations=FUNCTION_DECLARATIONS)
